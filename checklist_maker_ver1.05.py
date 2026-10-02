# -*- coding: utf-8 -*-
"""
照片 EXIF + 物种名录匹配工具（Tkinter 图形界面版 v6）

本版变化：
    - 只保留唯一一份已知名录："Checklist of plants of China 2026.xlsx"
      （香港和广东名录因格式不同已移除）
    - 未指定时，始终使用该默认名录
    - 自动流程不再弹选择框；仅保留"浏览…"按钮供手动指定

依赖：
    pip install Pillow openpyxl
    （可选）pip install pillow-heif
"""

import os
import queue
import re
import sys
import threading
import traceback
from datetime import datetime
from pathlib import Path

# ----------------- 依赖检查 -----------------
try:
    from PIL import Image
except ImportError:
    Image = None

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None

try:
    import pillow_heif
    pillow_heif.register_heif_opener()
except Exception:
    pass

import tkinter as tk
from tkinter import ttk, filedialog, messagebox


# ============================================================
# 常量配置
# ============================================================
# 唯一一份已知名录（默认使用）
DEFAULT_TEMPLATE_FILE = "Checklist of plants of China 2026.xlsx"

TEMPLATE_NAME = "物种名录模板.xlsx"          # 「生成模板」默认保存名
DEFAULT_OUTPUT_NAME = "照片信息"

# 物种名录必需的中文名称列
CN_NAME_COLUMN = "种中文名"

IMAGE_EXTENSIONS = {
    ".jpg", ".jpeg", ".jpe", ".jfif", ".png", ".bmp", ".gif",
    ".tif", ".tiff", ".webp", ".heic", ".heif", ".dng", ".cr2", ".nef",
}

TAG_DATETIME_ORIGINAL = 36867
TAG_DATETIME_DIGITIZED = 36868
TAG_DATETIME = 306
TAG_MAKE = 271
TAG_MODEL = 272
TAG_GPS_INFO = 34853

SPECIES_KEEP_COLS = [
    CN_NAME_COLUMN,
    "canonical_name", "genus", "genus_c", "species", "species_c",
    "infraspecies", "infraspecies_c", "infraspecies_marker",
    "author", "family", "family_c", "distribution_c", "id",
]
SPECIES_OUT_COLS = ["匹配_" + c for c in SPECIES_KEEP_COLS]

FIELDS = [
    "文件名", "完整路径", "文件大小(MB)",
    "拍摄日期", "拍摄时间", "拍摄日期时间",
    "纬度", "经度", "海拔(米)",
    "相机品牌", "相机型号",
    "图像宽度", "图像高度",
    "文件修改时间",
    "提取到的中文", "提取到的拉丁学名",
    *SPECIES_OUT_COLS,
    "备注",
]

NOISE_WORDS = {
    "img", "image", "photo", "photos", "pic", "pict", "picture", "pictures",
    "dsc", "dscn", "dscf", "copy", "final", "edited", "edit", "resized",
    "small", "large", "screenshot", "screen", "shot", "camera", "pano",
    "panorama", "the", "and", "for", "with", "from", "new", "old", "test",
    "jpeg", "jfif", "tiff", "heic", "heif", "webp",
}


# ============================================================
# 自定义异常
# ============================================================
class MissingSpeciesColumnError(Exception):
    """物种名录缺少必需列（例如「种中文名」）时抛出。"""
    pass


# ============================================================
# 通用工具
# ============================================================
def _as_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bytes):
        return value.decode("utf-8", "ignore").strip("\x00").strip()
    return str(value).strip()


def _dms_to_degrees(value):
    try:
        parts = list(value)
    except TypeError:
        return None
    if len(parts) < 3:
        return None
    try:
        d, m, s = float(parts[0]), float(parts[1]), float(parts[2])
    except (TypeError, ValueError):
        return None
    return d + m / 60.0 + s / 3600.0


def _parse_gps(gps):
    def get(tag_id, tag_name):
        if tag_id in gps:
            return gps[tag_id]
        if tag_name in gps:
            return gps[tag_name]
        return None

    lat = lon = alt = None
    lat_val = get(2, "GPSLatitude")
    lon_val = get(4, "GPSLongitude")
    lat_ref = _as_text(get(1, "GPSLatitudeRef")).upper()
    lon_ref = _as_text(get(3, "GPSLongitudeRef")).upper()
    alt_val = get(6, "GPSAltitude")
    alt_ref = get(5, "GPSAltitudeRef")

    if lat_val is not None:
        lat = _dms_to_degrees(lat_val)
        if lat is not None and lat_ref.startswith("S"):
            lat = -lat
    if lon_val is not None:
        lon = _dms_to_degrees(lon_val)
        if lon is not None and lon_ref.startswith("W"):
            lon = -lon
    if alt_val is not None:
        try:
            alt = float(alt_val)
            below = False
            if isinstance(alt_ref, bytes):
                below = len(alt_ref) > 0 and alt_ref[0] == 1
            elif alt_ref is not None:
                try:
                    below = int(alt_ref) == 1
                except (TypeError, ValueError):
                    below = False
            if below:
                alt = -alt
        except (TypeError, ValueError):
            alt = None
    return lat, lon, alt


def get_exe_dir() -> Path:
    """返回 exe（或 .py 脚本）所在目录。"""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


# ============================================================
# 物种名录搜索：只查找唯一的默认名录
# ============================================================
def _search_folders(photo_folder: Path = None) -> list:
    """返回搜索名录文件的文件夹列表（按优先级）。"""
    folders = []
    if photo_folder:
        folders.append(Path(photo_folder))
    folders.append(get_exe_dir())
    try:
        folders.append(Path.cwd())
    except Exception:
        pass
    for sub in ("Desktop", "Documents",
                "OneDrive/Desktop", "OneDrive/Documents"):
        try:
            folders.append(Path.home() / sub)
        except Exception:
            pass
    return folders


def find_default_template(photo_folder: Path = None, log=None):
    """
    在常见位置查找默认名录（大小写不敏感）。
    找到返回 Path；找不到返回 None。
    """
    folders = _search_folders(photo_folder)
    target_lower = DEFAULT_TEMPLATE_FILE.lower()
    for folder in folders:
        try:
            # 先试精确大小写
            p_exact = folder / DEFAULT_TEMPLATE_FILE
            if p_exact.exists() and p_exact.is_file():
                if log:
                    log(f"  默认名录已找到：{p_exact}")
                return p_exact
            # 再试大小写不敏感
            if not folder.is_dir():
                continue
            for f in folder.iterdir():
                if f.is_file() and f.name.lower() == target_lower:
                    if log:
                        log(f"  默认名录已找到：{f}")
                    return f
        except Exception:
            continue
    if log:
        log(f"  未找到默认名录「{DEFAULT_TEMPLATE_FILE}」")
    return None


def create_species_template(path: Path):
    """
    生成一个空白物种名录模板（带 2 行示例）。
    第一列即为必需的「种中文名」列。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "物种名录"

    ws.append(SPECIES_KEEP_COLS)

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="4472C4")
    for cell in ws[1]:
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    samples = [
        ["月季", "Rosa chinensis", "Rosa", "蔷薇属", "chinensis", "月季",
         "", "", "", "Jacq.", "Rosaceae", "蔷薇科",
         "中国各地普遍栽培", "示例-1"],
        ["银杏", "Ginkgo biloba", "Ginkgo", "银杏属", "biloba", "银杏",
         "", "", "", "L.", "Ginkgoaceae", "银杏科",
         "中国特产，浙江天目山有野生", "示例-2"],
    ]
    for s in samples:
        ws.append(s)

    for i, name in enumerate(SPECIES_KEEP_COLS, 1):
        w = max(len(str(name)) + 4, 14)
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    wb.save(path)


# ============================================================
# 文件名解析
# ============================================================
_DATE_PATTERNS = [
    re.compile(r"\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日?"),
    re.compile(r"\d{1,2}\s*月\s*\d{1,2}\s*日"),
    re.compile(r"\d{4}[-/.]\d{1,2}[-/.]\d{1,2}"),
]
_CN_CHUNK = re.compile(r"[\u4e00-\u9fff]+")
_LATIN_TOKEN = re.compile(r"[A-Za-z]+")


def extract_chinese_from_filename(stem: str):
    text = stem
    for pat in _DATE_PATTERNS:
        text = pat.sub(" ", text)
    chunks = _CN_CHUNK.findall(text)
    return "".join(chunks), chunks


def _looks_like_latin_pair(w1: str, w2: str) -> bool:
    if len(w1) < 3 or len(w2) < 3:
        return False
    if not (w1.isalpha() and w2.isalpha()):
        return False
    if not (w1[0].isupper() and w1[1:].islower()):
        return False
    if not w2.islower():
        return False
    if w1.lower() in NOISE_WORDS or w2.lower() in NOISE_WORDS:
        return False
    return True


def extract_latin_from_filename(stem: str, by_latin: dict):
    text = _CN_CHUNK.sub(" ", stem)
    tokens = _LATIN_TOKEN.findall(text)
    if by_latin:
        for i in range(len(tokens) - 1):
            pair = (tokens[i] + " " + tokens[i + 1]).lower()
            if pair in by_latin:
                return tokens[i] + " " + tokens[i + 1], by_latin[pair]
    for i in range(len(tokens) - 1):
        w1, w2 = tokens[i], tokens[i + 1]
        if _looks_like_latin_pair(w1, w2):
            return w1 + " " + w2, None
    return "", None


def analyze_filename(stem: str, by_cn: dict, by_latin: dict):
    cn_str, cn_chunks = extract_chinese_from_filename(stem)
    matched = None
    if cn_str:
        for chunk in cn_chunks:
            if chunk in by_cn:
                matched = by_cn[chunk]
                break
        if matched is not None:
            return cn_str, "", matched
    latin_str, latin_matched = extract_latin_from_filename(stem, by_latin)
    if matched is None and latin_matched is not None:
        matched = latin_matched
    return cn_str, latin_str, matched


# ============================================================
# 加载物种名录
# ============================================================
def load_species_db(db_path: Path, log=None):
    """
    读取物种名录。
    必须包含「种中文名」列，否则抛出 MissingSpeciesColumnError。
    返回 (by_cn, by_latin, keep_cols)
    """
    if not db_path or not Path(db_path).exists():
        if log:
            log(f"  名录文件不存在：{db_path}")
        return {}, {}, []

    db_path = Path(db_path)
    if log:
        log(f"  读取名录：{db_path}")

    wb = load_workbook(db_path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        wb.close()
        raise MissingSpeciesColumnError(
            f"物种名录文件是空的，无法读取表头：\n{db_path}"
        )

    header = [str(c).strip() if c is not None else "" for c in header_row]
    col_idx = {name: i for i, name in enumerate(header) if name}

    if CN_NAME_COLUMN not in col_idx:
        wb.close()
        raise MissingSpeciesColumnError(
            f"物种名录缺少必需列「{CN_NAME_COLUMN}」。\n\n"
            f"文件：{db_path}\n"
            f"现有表头：{header}\n\n"
            f"请在 Excel 中为该文件添加一列，列名必须是「{CN_NAME_COLUMN}」，\n"
            f"并在每一行填写物种的中文名称。\n\n"
            f"或者：点击主界面上的「生成模板」按钮，创建一个包含该列的空白模板，\n"
            f"然后把你的数据粘贴进去。"
        )

    keep = [c for c in SPECIES_KEEP_COLS if c in col_idx]

    by_cn: dict = {}
    by_latin: dict = {}

    for row in rows:
        if row is None:
            continue
        rec = {}
        for c in keep:
            idx = col_idx[c]
            v = row[idx] if idx < len(row) else None
            rec[c] = v if v is not None else ""

        cn = str(rec.get(CN_NAME_COLUMN, "")).strip()
        if cn:
            for part in re.split(r"[,，;；|/\s]+", cn):
                part = part.strip()
                if part:
                    by_cn.setdefault(part, rec)

        latin = str(rec.get("canonical_name", "")).strip()
        if latin:
            by_latin.setdefault(latin.lower(), rec)
            parts = latin.split()
            if len(parts) >= 2:
                by_latin.setdefault((parts[0] + " " + parts[1]).lower(), rec)

        genus = str(rec.get("genus", "")).strip()
        sp = str(rec.get("species", "")).strip()
        if genus and sp:
            by_latin.setdefault((genus + " " + sp).lower(), rec)

    wb.close()
    return by_cn, by_latin, keep


# ============================================================
# 读取单张照片 EXIF
# ============================================================
def extract_photo_info(path: Path) -> dict:
    stat = path.stat()
    info = {
        "文件名": path.name,
        "完整路径": str(path.resolve()),
        "文件大小(MB)": round(stat.st_size / 1024 / 1024, 3),
        "拍摄日期": "", "拍摄时间": "", "拍摄日期时间": "",
        "纬度": None, "经度": None, "海拔(米)": None,
        "相机品牌": "", "相机型号": "",
        "图像宽度": None, "图像高度": None,
        "文件修改时间": datetime.fromtimestamp(stat.st_mtime),
        "提取到的中文": "", "提取到的拉丁学名": "",
        "备注": "",
    }

    try:
        with Image.open(path) as img:
            info["图像宽度"], info["图像高度"] = img.size
            try:
                exif = img.getexif()
            except Exception as exc:
                exif = None
                info["备注"] = f"EXIF 读取失败：{exc}"

            if not exif:
                if not info["备注"]:
                    info["备注"] = "没有 EXIF 信息"
                return info

            dt_raw = (exif.get(TAG_DATETIME_ORIGINAL)
                      or exif.get(TAG_DATETIME_DIGITIZED)
                      or exif.get(TAG_DATETIME))
            dt_text = _as_text(dt_raw)
            if dt_text:
                dt = None
                for fmt in ("%Y:%m:%d %H:%M:%S", "%Y-%m-%d %H:%M:%S",
                            "%Y:%m:%d %H:%M:%S.%f", "%Y-%m-%d %H:%M:%S.%f"):
                    try:
                        dt = datetime.strptime(dt_text, fmt)
                        break
                    except ValueError:
                        continue
                if dt is not None:
                    info["拍摄日期"] = dt.strftime("%Y-%m-%d")
                    info["拍摄时间"] = dt.strftime("%H:%M:%S")
                    info["拍摄日期时间"] = dt
                else:
                    info["拍摄日期"] = dt_text
            else:
                info["备注"] = "无 EXIF 拍摄时间"

            info["相机品牌"] = _as_text(exif.get(TAG_MAKE))
            info["相机型号"] = _as_text(exif.get(TAG_MODEL))

            gps = None
            try:
                gps = exif.get_ifd(TAG_GPS_INFO)
            except Exception:
                gps = None
            if not gps:
                try:
                    raw = img._getexif() or {}
                    gps = raw.get(TAG_GPS_INFO)
                except Exception:
                    gps = None
            if gps:
                lat, lon, alt = _parse_gps(gps)
                if lat is not None:
                    info["纬度"] = round(lat, 6)
                if lon is not None:
                    info["经度"] = round(lon, 6)
                if alt is not None:
                    info["海拔(米)"] = round(alt, 2)
    except Exception as exc:
        info["备注"] = f"打开失败：{exc}"

    return info


# ============================================================
# 写出 Excel
# ============================================================
def write_excel(infos, out_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "照片信息"

    headers = ["序号"] + FIELDS
    ws.append(headers)
    for idx, info in enumerate(infos, start=1):
        row = [idx] + [info.get(f, "") for f in FIELDS]
        row = ["" if v is None else v for v in row]
        ws.append(row)

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="4472C4")
    for cell in ws[1]:
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    col = {name: i + 1 for i, name in enumerate(headers)}
    for r in range(2, ws.max_row + 1):
        ws.cell(r, col["纬度"]).number_format = "0.000000"
        ws.cell(r, col["经度"]).number_format = "0.000000"
        ws.cell(r, col["海拔(米)"]).number_format = "0.00"
        for name in ("拍摄日期时间", "文件修改时间"):
            c = ws.cell(r, col[name])
            if isinstance(c.value, datetime):
                c.number_format = "yyyy-mm-dd hh:mm:ss"

    for col_idx, name in enumerate(headers, start=1):
        max_w = sum(2 if ord(ch) > 127 else 1 for ch in str(name))
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, col_idx).value
            if v is None:
                continue
            s = v.strftime("%Y-%m-%d %H:%M:%S") if isinstance(v, datetime) else str(v)
            w = sum(2 if ord(ch) > 127 else 1 for ch in s)
            if w > max_w:
                max_w = w
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_w + 3, 70)

    wb.save(out_path)


# ============================================================
# 主处理流程（子线程）
# ============================================================
def run_processing(folder: Path, db_path, output_name: str, recursive: bool,
                   log_func, progress_func, done_func):
    try:
        if Image is None:
            raise RuntimeError("缺少 Pillow 库，请先执行：pip install Pillow")
        if Workbook is None:
            raise RuntimeError("缺少 openpyxl 库，请先执行：pip install openpyxl")

        log_func(f"扫描文件夹：{folder}")

        # 1) 名录：调用方优先；没给就用默认名录
        by_cn, by_latin, species_cols = {}, {}, []
        actual_db = None

        if db_path and Path(db_path).exists():
            actual_db = Path(db_path)
        else:
            actual_db = find_default_template(folder, log_func)

        if actual_db:
            log_func(f"使用物种名录：{actual_db}")
            try:
                by_cn, by_latin, species_cols = load_species_db(actual_db, log_func)
            except MissingSpeciesColumnError as exc:
                log_func("✘ 物种名录格式错误！")
                log_func(str(exc))
                done_func(False, str(exc))
                return

            log_func(f"  名录已加载：中文名 {len(by_cn)} 条，拉丁学名 {len(by_latin)} 条")
            missing = [c for c in SPECIES_KEEP_COLS if c not in species_cols]
            if missing:
                log_func(f"  提示：名录中缺少这些可选列 → {missing}")
        else:
            log_func("  未找到默认名录，物种匹配将被跳过（照片仍会处理）")

        # 2) 找图片
        pattern = folder.rglob("*") if recursive else folder.glob("*")
        files = sorted(
            (p for p in pattern
             if p.is_file()
             and p.suffix.lower() in IMAGE_EXTENSIONS
             and not p.name.startswith("~$")),
            key=lambda p: str(p).lower(),
        )
        if not files:
            log_func("没有找到任何图片文件。")
            done_func(True, "没有找到任何图片文件。")
            return

        log_func(f"共找到 {len(files)} 张图片，开始处理 ...")
        progress_func(0, len(files))

        # 3) 逐张处理
        infos = []
        matched_count = 0
        for i, p in enumerate(files, 1):
            info = extract_photo_info(p)
            try:
                cn_str, latin_str, matched = analyze_filename(p.stem, by_cn, by_latin)
            except Exception as exc:
                cn_str, latin_str, matched = "", "", None
                info["备注"] = (info.get("备注", "") + f"；匹配异常：{exc}").strip("；")
            info["提取到的中文"] = cn_str
            info["提取到的拉丁学名"] = latin_str
            for col in SPECIES_KEEP_COLS:
                info["匹配_" + col] = matched.get(col, "") if matched else ""
            if matched:
                matched_count += 1
            infos.append(info)

            progress_func(i, len(files))
            if i % 20 == 0 or i == len(files):
                log_func(f"  已处理 {i}/{len(files)}")

        # 4) 输出 Excel
        out_name = output_name.strip() or DEFAULT_OUTPUT_NAME
        if not out_name.lower().endswith(".xlsx"):
            out_name += ".xlsx"
        out_path = folder / out_name

        log_func(f"正在写出 Excel：{out_path.name} ...")
        try:
            write_excel(infos, out_path)
        except PermissionError:
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            out_path = folder / f"{Path(out_name).stem}_{stamp}.xlsx"
            log_func(f"  目标文件被占用，改用：{out_path.name}")
            write_excel(infos, out_path)

        # 5) 汇总
        with_time = sum(1 for x in infos if x["拍摄日期"])
        with_gps = sum(1 for x in infos if x["纬度"] is not None)
        with_cn = sum(1 for x in infos if x["提取到的中文"])
        with_latin = sum(1 for x in infos if x["提取到的拉丁学名"])

        log_func("─" * 40)
        log_func("处理完成！")
        log_func(f"  图片总数        ：{len(infos)}")
        log_func(f"  含拍摄时间      ：{with_time}")
        log_func(f"  含 GPS 信息     ：{with_gps}")
        log_func(f"  提取到中文      ：{with_cn}")
        log_func(f"  提取到拉丁学名  ：{with_latin}")
        log_func(f"  物种名录匹配成功：{matched_count}")
        log_func(f"  结果已保存      ：{out_path}")

        done_func(True, str(out_path))

    except Exception as exc:
        log_func("发生错误：")
        log_func(traceback.format_exc())
        done_func(False, str(exc))


# ============================================================
# Tkinter 主界面
# ============================================================
class PhotoExifApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("照片 EXIF + 物种匹配工具 v6")
        self.geometry("920x720")
        self.minsize(800, 600)

        self.msg_queue = queue.Queue()
        self.worker = None

        self._build_ui()
        self.after(80, self._poll_queue)
        self.after(200, self._initial_db_search)

    # ---------------- UI 构建 ----------------
    def _build_ui(self):
        pad = {"padx": 10, "pady": 6}

        # 1. 照片文件夹
        frame_top = ttk.LabelFrame(self, text="1. 选择照片文件夹")
        frame_top.pack(fill="x", **pad)

        self.var_folder = tk.StringVar()
        ttk.Entry(frame_top, textvariable=self.var_folder).pack(
            side="left", fill="x", expand=True, padx=(10, 5), pady=10)
        ttk.Button(frame_top, text="浏览…", command=self._choose_folder).pack(
            side="left", padx=(0, 10), pady=10)

        # 2. 物种名录
        frame_db = ttk.LabelFrame(self, text="2. 物种名录文件（必须包含「种中文名」列）")
        frame_db.pack(fill="x", **pad)

        row1 = ttk.Frame(frame_db)
        row1.pack(fill="x", padx=10, pady=(8, 2))
        self.var_db = tk.StringVar()
        ttk.Entry(row1, textvariable=self.var_db).pack(
            side="left", fill="x", expand=True, padx=(0, 5))
        ttk.Button(row1, text="浏览…", width=8,
                   command=self._choose_db).pack(side="left", padx=2)
        ttk.Button(row1, text="自动查找", width=9,
                   command=self._auto_find_db).pack(side="left", padx=2)
        ttk.Button(row1, text="生成模板", width=9,
                   command=self._create_template).pack(side="left", padx=2)

        row2 = ttk.Frame(frame_db)
        row2.pack(fill="x", padx=10, pady=(0, 8))
        self.lbl_db_status = ttk.Label(
            row2, text="状态：未指定", foreground="gray")
        self.lbl_db_status.pack(side="left")

        ttk.Label(
            frame_db,
            text=(f"※ 名录中必须有一列名为「{CN_NAME_COLUMN}」。\n"
                  f"   默认（也是唯一的已知名录）：{DEFAULT_TEMPLATE_FILE}\n"
                  f"   程序会按以下顺序自动查找：照片文件夹 → 程序所在目录 →\n"
                  f"   当前工作目录 → 桌面 → 文档。\n"
                  f"   如需使用其它名录，请点击「浏览…」手动指定。"),
            foreground="#666", justify="left"
        ).pack(fill="x", padx=10, pady=(0, 8))

        # 3. 输出文件名
        frame_mid = ttk.LabelFrame(self, text="3. 输出 Excel 文件名")
        frame_mid.pack(fill="x", **pad)

        self.var_outname = tk.StringVar(value=DEFAULT_OUTPUT_NAME)
        ttk.Entry(frame_mid, textvariable=self.var_outname).pack(
            side="left", fill="x", expand=True, padx=(10, 5), pady=10)
        ttk.Label(frame_mid, text=".xlsx").pack(side="left", padx=(0, 10))

        # 4. 选项
        frame_opt = ttk.LabelFrame(self, text="4. 选项")
        frame_opt.pack(fill="x", **pad)

        self.var_recursive = tk.BooleanVar(value=True)
        ttk.Checkbutton(frame_opt, text="包含子文件夹",
                        variable=self.var_recursive).pack(
            side="left", padx=12, pady=8)

        # 按钮
        frame_btn = ttk.Frame(self)
        frame_btn.pack(fill="x", **pad)

        self.btn_start = ttk.Button(frame_btn, text="开始处理",
                                    command=self._on_start)
        self.btn_start.pack(side="left", padx=(10, 6), pady=4)

        self.btn_open = ttk.Button(frame_btn, text="打开输出文件夹",
                                   command=self._open_output_folder,
                                   state="disabled")
        self.btn_open.pack(side="left", padx=6, pady=4)

        self.btn_clear = ttk.Button(frame_btn, text="清空日志",
                                    command=self._clear_log)
        self.btn_clear.pack(side="left", padx=6, pady=4)

        # 进度条
        frame_prog = ttk.Frame(self)
        frame_prog.pack(fill="x", **pad)

        self.var_prog = tk.StringVar(value="就绪")
        ttk.Label(frame_prog, textvariable=self.var_prog, width=22).pack(
            side="left", padx=(10, 6))
        self.progress = ttk.Progressbar(frame_prog, mode="determinate",
                                        maximum=100)
        self.progress.pack(side="left", fill="x", expand=True,
                           padx=(0, 10), pady=4)

        # 日志
        frame_log = ttk.LabelFrame(self, text="处理日志")
        frame_log.pack(fill="both", expand=True, **pad)

        self.txt_log = tk.Text(frame_log, wrap="word", height=14,
                               font=("Consolas", 10))
        scroll = ttk.Scrollbar(frame_log, orient="vertical",
                               command=self.txt_log.yview)
        self.txt_log.configure(yscrollcommand=scroll.set)
        scroll.pack(side="right", fill="y")
        self.txt_log.pack(side="left", fill="both", expand=True,
                          padx=(8, 0), pady=8)
        self.txt_log.configure(state="disabled")

    # ---------------- 名录相关 ----------------
    def _set_db_status(self, path: Path):
        if path and Path(path).exists():
            self.var_db.set(str(path))
            name = Path(path).name
            if name.lower() == DEFAULT_TEMPLATE_FILE.lower():
                self.lbl_db_status.configure(
                    text="状态：默认名录已就绪 ✔", foreground="#0a8a0a")
            else:
                self.lbl_db_status.configure(
                    text=f"状态：使用自定义名录（{name}）",
                    foreground="#0a6aa0")
        else:
            self.lbl_db_status.configure(
                text="状态：未找到名录 ⚠", foreground="#c04000")

    def _initial_db_search(self):
        default_path = find_default_template(None, self._append_log)
        if default_path:
            self._set_db_status(default_path)
            self._append_log(f"启动时使用默认名录：{default_path}")
        else:
            self.lbl_db_status.configure(
                text="状态：未找到默认名录 ⚠", foreground="#c04000")
            self._append_log(
                f"未在常见位置找到默认名录「{DEFAULT_TEMPLATE_FILE}」。\n"
                f"您可以：① 点击「浏览…」手动指定；② 点击「生成模板」创建空白模板；\n"
                f"③ 直接点「开始处理」（照片仍会处理，但不匹配物种信息）。")

    def _choose_db(self):
        path = filedialog.askopenfilename(
            title="选择物种名录 Excel 文件",
            filetypes=[("Excel 文件", "*.xlsx *.xls"), ("所有文件", "*.*")])
        if path:
            self._set_db_status(Path(path))

    def _auto_find_db(self):
        folder = self.var_folder.get().strip()
        photo_folder = Path(folder) if folder else None

        default_path = find_default_template(photo_folder, self._append_log)
        if default_path:
            self._set_db_status(default_path)
            messagebox.showinfo(
                "已找到",
                f"已找到默认名录：\n{default_path}")
        else:
            messagebox.showwarning(
                "未找到",
                f"未在常见位置找到默认名录：\n{DEFAULT_TEMPLATE_FILE}\n\n"
                f"您可以点击「浏览…」手动指定，或点击「生成模板」创建空白模板。")

    def _create_template(self):
        folder = self.var_folder.get().strip()
        base = Path(folder) if folder and Path(folder).is_dir() else get_exe_dir()
        path = filedialog.asksaveasfilename(
            title="保存空白物种名录模板",
            initialdir=str(base),
            initialfile=TEMPLATE_NAME,
            defaultextension=".xlsx",
            filetypes=[("Excel 文件", "*.xlsx")])
        if not path:
            return
        try:
            create_species_template(Path(path))
        except Exception as exc:
            messagebox.showerror("生成失败", str(exc))
            return
        self._set_db_status(Path(path))
        messagebox.showinfo(
            "模板已生成",
            f"已在以下位置生成模板：\n{path}\n\n"
            f"模板第一列就是必需的「{CN_NAME_COLUMN}」列。\n"
            f"请用 Excel 打开，往里面填入中文名、拉丁名、属名、科名等信息后，\n"
            f"再点击「开始处理」即可参与匹配。")

    # ---------------- 其他事件 ----------------
    def _choose_folder(self):
        folder = filedialog.askdirectory(title="选择照片所在的文件夹")
        if folder:
            self.var_folder.set(folder)
            # 若尚未指定名录，尝试用默认名录
            if not self.var_db.get().strip():
                default_path = find_default_template(Path(folder))
                if default_path:
                    self._set_db_status(default_path)

    def _open_output_folder(self):
        folder = self.var_folder.get().strip()
        if not folder:
            return
        path = Path(folder)
        if not path.is_dir():
            messagebox.showerror("错误", f"文件夹不存在：{folder}")
            return
        try:
            if sys.platform.startswith("win"):
                os.startfile(str(path))
            elif sys.platform == "darwin":
                import subprocess
                subprocess.Popen(["open", str(path)])
            else:
                import subprocess
                subprocess.Popen(["xdg-open", str(path)])
        except Exception as exc:
            messagebox.showerror("无法打开", str(exc))

    def _clear_log(self):
        self.txt_log.configure(state="normal")
        self.txt_log.delete("1.0", "end")
        self.txt_log.configure(state="disabled")

    def _on_start(self):
        if self.worker and self.worker.is_alive():
            messagebox.showinfo("提示", "任务正在运行中，请稍候……")
            return

        folder_str = self.var_folder.get().strip()
        if not folder_str:
            messagebox.showwarning("提示", "请先选择照片所在文件夹。")
            return
        folder = Path(folder_str).expanduser()
        if not folder.is_dir():
            messagebox.showerror("错误", f"文件夹不存在：{folder}")
            return

        # ---- 解析名录路径 ----
        db_path = None
        db_str = self.var_db.get().strip()

        if db_str:
            p = Path(db_str)
            if p.exists() and p.is_file():
                db_path = p
            else:
                if not messagebox.askyesno(
                        "名录不存在",
                        f"指定的名录文件不存在：\n{p}\n\n"
                        "是否仍要继续？（将无法匹配物种信息）"):
                    return

        if db_path is None:
            # 默认名录
            default_path = find_default_template(folder, self._append_log)
            if default_path:
                db_path = default_path
                self._set_db_status(default_path)
            else:
                ans = messagebox.askyesnocancel(
                    "未找到默认名录",
                    f"未找到默认名录：\n{DEFAULT_TEMPLATE_FILE}\n\n"
                    "• 点击【是】：继续处理（照片仍会处理，但不匹配物种信息）\n"
                    "• 点击【否】：先取消，去准备名录文件\n"
                    "• 点击【取消】：取消处理\n\n"
                    "提示：名录中必须包含一列「种中文名」。\n"
                    "      也可以点击「生成模板」快速创建一份空白名录。")
                if ans is None or ans is False:
                    return

        out_name = self.var_outname.get().strip() or DEFAULT_OUTPUT_NAME

        self._clear_log()
        self.btn_start.configure(state="disabled")
        self.btn_open.configure(state="disabled")
        self.progress["value"] = 0
        self.var_prog.set("处理中…")

        self.worker = threading.Thread(
            target=run_processing,
            args=(folder, db_path, out_name, self.var_recursive.get(),
                  self._log_from_thread,
                  self._progress_from_thread,
                  self._done_from_thread),
            daemon=True,
        )
        self.worker.start()

    # ---------------- 子线程 -> 主线程 ----------------
    def _log_from_thread(self, text: str):
        self.msg_queue.put(("log", text))

    def _progress_from_thread(self, current: int, total: int):
        self.msg_queue.put(("progress", current, total))

    def _done_from_thread(self, ok: bool, payload: str):
        self.msg_queue.put(("done", ok, payload))

    def _poll_queue(self):
        try:
            while True:
                msg = self.msg_queue.get_nowait()
                kind = msg[0]

                if kind == "log":
                    self._append_log(msg[1])
                elif kind == "progress":
                    cur, total = msg[1], msg[2]
                    if total > 0:
                        self.progress["value"] = cur / total * 100
                    self.var_prog.set(f"已处理 {cur}/{total}")
                elif kind == "done":
                    ok, payload = msg[1], msg[2]
                    self.btn_start.configure(state="normal")
                    if ok:
                        self.var_prog.set("完成")
                        self.progress["value"] = 100
                        self.btn_open.configure(state="normal")
                        messagebox.showinfo(
                            "完成",
                            f"处理完成！\n\n结果已保存到：\n{payload}")
                    else:
                        self.var_prog.set("出错")
                        messagebox.showerror("出错了", payload)
        except queue.Empty:
            pass
        finally:
            self.after(80, self._poll_queue)

    def _append_log(self, text: str):
        self.txt_log.configure(state="normal")
        self.txt_log.insert("end", text + "\n")
        self.txt_log.see("end")
        self.txt_log.configure(state="disabled")


# ============================================================
# 入口
# ============================================================
def main():
    if Image is None or Workbook is None:
        missing = []
        if Image is None:
            missing.append("Pillow")
        if Workbook is None:
            missing.append("openpyxl")
        root = tk.Tk()
        root.withdraw()
        messagebox.showerror(
            "缺少依赖",
            "缺少以下库，请先在命令行执行：\n\n"
            f"pip install {' '.join(missing)}\n")
        root.destroy()
        return

    app = PhotoExifApp()
    app.mainloop()


if __name__ == "__main__":
    main()